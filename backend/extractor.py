import io
import docx
import fitz
from pptx import Presentation

def extract_text(file_path_obj):
    # In FastAPI, we will receive a file object or path string.
    # Let's handle an UploadFile or a local path.
    # Assuming it's a temp file path saved by FastAPI or directly read from UploadFile bytes
    try:
        # If it's passed as a tuple from UploadFile, unpack it (filename, bytes)
        if isinstance(file_path_obj, tuple):
            filename, content = file_path_obj
            file_ext = "." + filename.split('.')[-1].lower() if '.' in filename else ""
        elif isinstance(file_path_obj, str):
            filename = file_path_obj
            file_ext = "." + filename.split('.')[-1].lower() if '.' in filename else ""
            with open(filename, 'rb') as f:
                content = f.read()
        else:
            return "Error: Unsupported input format"

        text = ""
        
        if file_ext == '.txt':
            text = content.decode('utf-8', errors='replace')
        elif file_ext == '.docx':
            doc_file = io.BytesIO(content)
            doc = docx.Document(doc_file)
            text = '\n'.join([para.text for para in doc.paragraphs])
        elif file_ext == '.pdf':
            doc_file = io.BytesIO(content)
            pdf_doc = fitz.open(stream=doc_file, filetype="pdf")
            for page in pdf_doc:
                text += page.get_text() + "\n"
        elif file_ext == '.pptx':
            doc_file = io.BytesIO(content)
            prs = Presentation(doc_file)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text + "\n"
        else:
            return "Error: Unsupported file format. Please upload .txt, .docx, .pdf, or .pptx."
            
        return text.strip()
    except Exception as e:
        return f"Error: Failed to process file - {str(e)}"
